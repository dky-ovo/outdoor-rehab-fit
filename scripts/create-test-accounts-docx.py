"""Create a private credential handover from confirmed account-creation receipts."""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


def generate(source, output):
    accounts = json.loads(source.read_text(encoding="utf-8"))
    if len(accounts) != 50 or not all(a.get("created") is True for a in accounts):
        raise ValueError("Exactly 50 confirmed account creations are required")
    if len({a["email"] for a in accounts}) != 50 or len({a["password"] for a in accounts}) != 50:
        raise ValueError("Duplicate credentials")
    doc = Document()
    doc.core_properties.title = "OUTDOOR-FIT — 50 個壓力測試帳戶"
    doc.core_properties.subject = "Private test account credentials"
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(14)
    section.left_margin = section.right_margin = Mm(13.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.paragraph_format.space_after = Pt(4)
    header = section.header.paragraphs[0]
    header.text = "OUTDOOR-FIT  |  私人帳戶資料 · 請勿公開"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string("840B1C")
    footer = section.footer.paragraphs[0]
    footer.text = "建立日期：2026-10-02  ·  只供壓力測試  |  "
    page = OxmlElement("w:fldSimple"); page.set(qn("w:instr"), "PAGE"); footer._p.append(page)
    widths = [10, 30, 86, 57]
    for page_index in range(2):
        if page_index:
            doc.add_page_break()
        heading = doc.add_paragraph()
        run = heading.add_run("50 個壓力測試帳戶" if not page_index else "壓力測試帳戶（續）")
        run.bold = True; run.font.size = Pt(19); run.font.color.rgb = RGBColor.from_string("840B1C")
        if not page_index:
            doc.add_paragraph("狀態：已在正式後台建立。類型：普通參與者，沒有後台管理權限。")
            doc.add_paragraph("登入：https://app.outdoorfit.hk/ 或手機 APP；選擇電郵登入，使用下表的完整登入名及密碼。")
            doc.add_paragraph("帳戶使用獨立隨機密碼；電郵與電話為測試佔位資料，不接收短信或郵件。測試資料會寫入正式環境，使用時請選取測試帳戶。")
        else:
            doc.add_paragraph("以下帳戶均已建立。電郵是登入名；「壓力測試」是後台顯示名稱。")
        table = doc.add_table(rows=1, cols=4)
        table.autofit = False
        table.style = "Table Grid"
        for column, width in zip(table.columns, widths):
            column.width = Mm(width)
        for cell, label, width in zip(table.rows[0].cells, ["序號", "顯示名稱", "登入名（電郵）", "密碼"], widths):
            cell.width = Mm(width); cell.text = label
            shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "840B1C"); cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs:
                run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
        repeat = OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for account in accounts[page_index * 25:(page_index + 1) * 25]:
            row = table.add_row()
            values = [str(account["number"]), account["name"], account["email"], account["password"]]
            for cell, value, width in zip(row.cells, values, widths):
                cell.width = Mm(width); cell.text = value; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
                for run in p.runs: run.font.size = Pt(9)
            no_split = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(no_split)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    check = Document(output)
    cells = [c.text for table in check.tables for row in table.rows[1:] for c in row.cells]
    for account in accounts:
        if account["email"] not in cells or account["password"] not in cells:
            raise ValueError("Credential roundtrip failed")
    print("DOCX_ACCOUNTS=50")
    print("DOCX_CONTENT_VERIFIED=PASS")
    print("DOCX_FILE=" + str(output.resolve()))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    generate(args.source, args.output)
