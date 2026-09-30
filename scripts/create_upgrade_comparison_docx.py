from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "output" / "OUTDOOR-FIT_App_Upgrade_Comparison_CN-EN.docx"

rows = [
    ("Login and guest access\n登录及访客权限", "○ No login, registration or guest-access control; all users can access the existing functions.\n无登录、注册或访客权限；所有用户可进入现有功能。", "Login, registration, guest access restrictions, login prompts, server-based accounts, password security and cross-device data synchronisation.\n登录、注册、访客权限限制、登录提示，以及服务器式账户、密码安全和跨设备资料同步。"),
    ("GIS equipment location\nGIS 器材定位", "◐ A Location page with an embedded web map is available.\n已有地点页面，以嵌入式网页地图显示。", "Add equipment filtering, nearby search, geolocation, routing and equipment-detail linkage once GIS data are available.\n待 GIS 数据提供后，加入器材类别筛选、附近搜寻、定位、路线导航及器材详情联动。"),
    ("Content Management System\n内容管理系统", "○ Content is embedded in local app code and assets.\n内容存于 App 本地代码及资源档案。", "Admin portal, editing, publishing, versioning, review and remote synchronisation.\n管理员后台、内容编辑、发布、版本控制、审核及远端同步。"),
    ("Usage data and automatic logout\n使用数据及自动登出", "○ No tracking of active days, page views, audio plays, goal achievement or login events; no automatic logout.\n未记录使用天数、页面浏览、音频播放、目标达成或登录事件；无自动登出。", "Event tracking, research database, consent/privacy controls, reporting export and automatic logout after 1.5 hours of inactivity.\n事件追踪、研究资料库、同意/隐私机制、报表导出，以及 1.5 小时无操作自动登出。"),
    ("Exercise session record\n运动记录", "○ No session timer, summary or session history.\n无运动计时、运动总结或历史记录。", "Support outdoor/home sessions with start, pause, resume and end controls; save time, duration, emotion, perceived exertion, journal, equipment and repetitions.\n支援户外及家居运动的开始、暂停、继续和结束计时；保存时间、时长、情绪、主观运动强度、日志、器材及次数。"),
    ("Photo-based equipment recognition\n器材拍照识别", "○ No camera or equipment-recognition function.\n无拍照或器材识别功能。", "Camera permission, capture/upload, image recognition, automatic exercise-name entry, and manual confirmation/correction.\n相机权限、拍摄/上传、器材图像识别、自动填入运动名称，以及人工确认和更正流程。"),
    ("Session review and sharing\n运动记录查看及分享", "○ No session-record function.\n无运动记录功能。", "Session log, individual record details, share button, summary card, and sharing to the group or external social platforms.\n运动日志、单次记录详情、分享按钮、分享摘要卡片，以及群组或外部社交平台分享。"),
    ("Goal setting and progress tracking\n目标设定及进度追踪", "○ No weekly exercise goals, SMART goals or progress display.\n无每周运动目标、SMART 目标或进度显示。", "Weekly aerobic, balance and strengthening targets; adjustable target days; SMART-goal guidance, calendar, automatic calculation and progress displays on the home/summary pages.\n有氧、平衡及肌力三类每周目标；可调整目标天数；SMART 目标指引、日历、自动统计及首页/总结页进度显示。"),
    ("Group messaging\n群组消息（SMS）", "○ No group communication feature.\n无群组沟通功能。", "Group text messages, photos, reactions, notifications, member management, privacy controls and content moderation.\n群组文字讯息、照片、表情反应、通知、成员管理、隐私控制及内容审核。"),
    ("Chatbot and reminders\n聊天机器人及提醒", "○ No chatbot or automated reminders.\n无聊天机器人或自动提醒。", "FAQ knowledge base, weekly reminder schedule, group-leader role, response boundaries and human-support process.\n常见问题知识库、每周提醒排程、群组负责人角色、回复边界及人工支援机制。"),
    ("Expanded equipment and video content\n扩充器材及影片内容", "✅ The App already includes an equipment catalogue, text explanations, exercise animations, bilingual pre-recorded audio and YouTube videos.\n已有器材目录、文字说明、动作动画、双语预录语音及 YouTube 影片。", "Expand content using equipment materials and YouTube links supplied by the client; CMS-based maintenance is recommended.\n按需求方后续提供的器材资料及 YouTube 链接扩充内容；建议透过 CMS 维护。"),
    ("Learning progress in My Favourite\n收藏学习进度", "◐ Equipment can be bookmarked, but reading, viewing and listening activity is not tracked.\n可收藏器材，但不记录是否已阅读、观看或收听。", "Reading, video and audio completion events; individual learning progress, reminders and research reports.\n阅读、影片及语音完成事件；个人学习进度、提醒及研究报表。"),
    ("Text-to-Speech\n文字转语音", "◐ Pre-recorded bilingual equipment audio is available, but not dynamic text-to-speech.\n现有每项器材的中英文预录语音，不是即时文字转语音。", "TTS engine, language/voice selection, playback controls, offline strategy and accessibility design.\nTTS 引擎、语言/声线选择、播放控制、离线策略及无障碍设计。"),
]

def set_font(run, name="Aptos", size=9, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_table_widths(table, widths):
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblLayout = OxmlElement("w:tblLayout")
    tblLayout.set(qn("w:type"), "fixed")
    tblPr.append(tblLayout)
    grid = table._tbl.tblGrid
    for i, width in enumerate(widths):
        grid.gridCol_lst[i].set(qn("w:w"), str(width))
    for row in table.rows:
        for i, width in enumerate(widths):
            tcW = row.cells[i]._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                row.cells[i]._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(width))
            tcW.set(qn("w:type"), "dxa")

def write_cell(cell, text, header=False):
    cell.text = ""
    for index, line in enumerate(text.split("\n")):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5 if not header else 0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(line)
        set_font(r, size=8.2 if not header else 8.6, bold=header, color=(255,255,255) if header else (25,35,45))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11)
section.page_height = Inches(8.5)
section.top_margin = Inches(0.48)
section.bottom_margin = Inches(0.48)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)
section.header_distance = Inches(0.25)
section.footer_distance = Inches(0.25)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(9)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.paragraph_format.space_after = Pt(0)
run = header.add_run("OUTDOOR-FIT APP | Upgrade Requirements Comparison")
set_font(run, size=7.5, color=(100,110,120))

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_before = Pt(0)
run = footer.add_run("Prepared for requirements confirmation | 需求确认文件")
set_font(run, size=7.5, color=(100,110,120))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("OUTDOOR-FIT App Upgrade: Current App and Required Scope")
set_font(r, size=17, bold=True, color=(24,72,108))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("户外健体乐 App 升级：现有 App 与正式升级需求对照")
set_font(r, size=12, bold=True, color=(55,85,110))

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
r = p.add_run("Status / 状态: ")
set_font(r, size=8.5, bold=True, color=(55,85,110))
r = p.add_run("✅ Implemented / 已实现     ◐ Partially implemented / 部分实现     ○ Not implemented / 未实现")
set_font(r, size=8.5, color=(60,60,60))

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
set_table_widths(table, [2050, 3050, 5160])
headers = ["Upgrade Requirement\n升级需求", "Current App Status\n现有 App 功能", "Required for Production Upgrade\n正式升级所需"]
for cell, text in zip(table.rows[0].cells, headers):
    shade(cell, "1F4E78")
    write_cell(cell, text, header=True)
set_repeat_table_header(table.rows[0])

for idx, row_data in enumerate(rows):
    cells = table.add_row().cells
    fill = "F5F8FB" if idx % 2 == 0 else "FFFFFF"
    for cell, text in zip(cells, row_data):
        shade(cell, fill)
        write_cell(cell, text)

doc.add_paragraph().paragraph_format.space_after = Pt(0)
heading = doc.add_paragraph()
heading.paragraph_format.space_before = Pt(7)
heading.paragraph_format.space_after = Pt(2)
r = heading.add_run("Suggested Client Wording / 对外说明建议")
set_font(r, size=10, bold=True, color=(31,78,120))

for text in [
    "The current App provides core functions including outdoor-equipment learning content, bilingual information, pre-recorded audio, videos, favourites, weather information, risk assessment, safety guidance and a location map. The listed upgrades will be developed and accepted in phases, subject to final requirements, GIS/content-data availability, backend architecture, and research privacy/data-management requirements.",
    "现有 App 已提供户外器材教学内容、双语资讯、预录语音、影片、收藏、天气资讯、风险评估、安全说明及地点地图等基础功能。上述升级项目将按正式需求、GIS 及内容资料供应、后端架构，以及研究私隐和数据管理要求分阶段开发及验收。",
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_font(r, size=8.5, color=(45,45,45))

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
