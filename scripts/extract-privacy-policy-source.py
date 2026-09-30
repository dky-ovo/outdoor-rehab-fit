"""Read the approved DOCX into a deterministic test fixture; never change the DOCX."""
import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from lxml import etree

SOURCE_SHA256 = "418E64A40766E1F5D16BE6F571514250799A3E56953565DC044FD700DF596E83"
parser = argparse.ArgumentParser()
parser.add_argument("source")
parser.add_argument("output")
args = parser.parse_args()
source = Path(args.source)
digest = hashlib.sha256(source.read_bytes()).hexdigest().upper()
if digest != SOURCE_SHA256:
    raise ValueError("Unexpected source revision; review the changed policy before importing it")
document = Document(source)
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
with ZipFile(source) as archive:
    xml = etree.fromstring(archive.read("word/document.xml"))
    changes = len(xml.findall(".//w:ins", ns)) + len(xml.findall(".//w:del", ns))
    comments = etree.fromstring(archive.read("word/comments.xml"))
    comment_count = len(comments.findall("w:comment", ns))
    media_count = sum(name.startswith("word/media/") for name in archive.namelist())
if changes or document.tables or media_count:
    raise ValueError("Unexpected tracked changes, tables or media require manual review")
fixture = {
    "sourceSha256": digest,
    "trackedChanges": changes,
    "excludedReviewerComments": comment_count,
    "blocks": [{"index": i, "style": p.style.name, "text": p.text} for i, p in enumerate(document.paragraphs)],
}
output = Path(args.output)
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps(fixture, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"blocks": len(fixture["blocks"]), "excludedReviewerComments": comment_count, "sourceSha256": digest}))
